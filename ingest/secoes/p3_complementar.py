# -*- coding: utf-8 -*-
r"""
Ingestão P3 — carga COMPLEMENTAR (não mexe nos indicadores RAC já cobertos
por ingest/secoes/p3_rac.py: abordados_carros, veiculos_recuperados,
flagrantes, capturas_procurado etc.).

Fontes reais (achadas por varredura em 2026-07-19, confirmadas via
manifest scratchpad/staging/manifests/P3.jsonl):

1) \\cmdo\pmesp\16BPMM\16BPMM_EM\P3\Documentos P3\2026\ESTATÍSTICAS\PRODUTIVIDADE\
   - "PRODUTIVIDADE PARAISÓPOLIS 2026.xlsx" (aba PRODUTIVIDADE, 1 linha por
     evento/dia da Operação Paz e Proteção Paraisópolis) — soma mensal da
     coluna PESSOAS ABORDADAS -> indicador novo 'produtividade_paraisopolis'.
   - "RESULTADOS OPERAÇÕES RPM 2026.xlsx" (1 aba por mês JAN26..AGO26, layout
     matriz: linha=métrica, coluna=dia do mês, última coluna=TOTAL) — pega o
     TOTAL da linha "Pessoas abordadas" de cada mês -> indicador novo
     'resultado_operacoes_rpm'. Meses com TOTAL=0 (ex.: AGO26 — aba com dados
     de 01-31JUL26 no cabeçalho mas todas as células vazias, sobra de
     copiar/colar do template) são descartados por não representarem dado
     real.
   - "PRODUTIVIDADE OPERAÇÕES.xlsx" — NÃO carregado nesta rodada: é o log
     bruto (1 linha por evento) de TODAS as operações (Força Tática, RPM,
     CIAs, PSIU etc. juntas, com coluna CIAS/OPERAÇÃO), sobreposto aos dois
     arquivos acima; carregar geraria dupla contagem sem que a tarefa tenha
     definido uma terceira chave de indicador dedicada. Fica registrado como
     não carregado (ver retorno da tarefa).

2) \\cmdo\pmesp\16BPMM\16BPMM_EM\P3\Documentos P3\2026\ESTATÍSTICAS\
   INDICADORES CRIMINAIS PARA GOOGLE MAPS\<NNMESAA>\*.xlsx — um arquivo por
   natureza-mês (ex.: "01JAN26 ROUBO DE VEÍCULO.xlsx"), aba "MAPS" traz 1
   linha por BO (deduplicado; a aba "Sheet1"/"Plan1" tem 1 linha por VERSÃO
   do BO e infla a contagem). Alguns meses têm arquivos duplicados/quebrados
   em I/II/III (achado real: em fev/mar26 as variantes sem sufixo e I/II
   vieram vazias — só a última parte tem os dados do mês inteiro); por isso
   a extração pega a UNIÃO dos BOs de todas as variantes de uma mesma
   categoria-mês, não a soma bruta de linhas. Mapeado para os indicadores
   novos 'roubo_registrado', 'furto_registrado', 'homicidio_registrado',
   'roubo_carga_registrado' (mensal, cia=0 — os arquivos não têm quebra por
   Cia). Só há pastas de 01JAN26 a 06JUN26 nesta rodada (jul/26 ainda não
   publicado nesta pasta).

3) Snapshots locais do painel Muralha (rolinha de 3 dias, ver aba "Notas" do
   próprio xlsx — NÃO é histórico, é o acumulado corrente da tarefa
   agendada do CCO-16):
   - C:\Users\13934852785\Desktop\CCO-16_Ocorrencias.xlsx (aba "Ocorrencias")
   - C:\Users\13934852785\Downloads\cco16_backup.csv
   Os dois arquivos são o MESMO feed (Muralha Paulista, áreas do 16º BPM/M)
   capturado em momentos próximos — deduplicados por
   (Data/Hora Evento, Tipo, Endereço) antes de contar, senão o mesmo evento
   entraria duas vezes. Carregado em agregado_dimensional
   (secao='p3', fonte='ocorrencias_muralha', dimensao='natureza'), contagem
   por natureza extraída do campo Tipo (prefixo "<código> - " removido;
   "SSP:"/"LAP:" tratados à parte). É um retrato pontual (gerado
   19/07/2026), não uma série — não confundir com os indicadores mensais
   acima.

Uso:
    python -m ingest.secoes.p3_complementar [--dry-run] [--forcar]
"""
import argparse
import csv
import glob
import os
import re
import sys
from datetime import datetime, timezone

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
from ingest.common.unc import caminho_disponivel
from ingest.common.excel import mtime_iso, numero_ou_none
from ingest.common.carga import (
    abrir_batch,
    fechar_batch,
    upsert_fato_secao,
    upsert_agregado_dimensional,
    registrar_arquivo,
    arquivo_ja_ingerido,
    sha256_arquivo,
)

SECAO = "p3"

RAIZ_PRODUTIVIDADE = r"\\cmdo\pmesp\16BPMM\16BPMM_EM\P3\Documentos P3\2026\ESTATÍSTICAS\PRODUTIVIDADE"
RAIZ_GOOGLE_MAPS = r"\\cmdo\pmesp\16BPMM\16BPMM_EM\P3\Documentos P3\2026\ESTATÍSTICAS\INDICADORES CRIMINAIS PARA GOOGLE MAPS"
ARQUIVO_PARAISOPOLIS = os.path.join(RAIZ_PRODUTIVIDADE, "PRODUTIVIDADE PARAISÓPOLIS 2026.xlsx")
ARQUIVO_RPM = os.path.join(RAIZ_PRODUTIVIDADE, "RESULTADOS OPERAÇÕES RPM 2026.xlsx")

OCORRENCIAS_XLSX = r"C:\Users\13934852785\Desktop\CCO-16_Ocorrencias.xlsx"
OCORRENCIAS_CSV = r"C:\Users\13934852785\Downloads\cco16_backup.csv"

MESES_ABREV = {
    "JAN": 1, "FEV": 2, "MAR": 3, "ABR": 4, "MAI": 5, "JUN": 6,
    "JUL": 7, "AGO": 8, "SET": 9, "OUT": 10, "NOV": 11, "DEZ": 12,
}

CATEGORIA_INDICADOR = {
    "ROUBO DE CARGA": "roubo_carga_registrado",
    "ROUBO CARGA": "roubo_carga_registrado",
    "ROUBO DE VEICULO": "roubo_registrado",
    "ROUBO OUTROS": "roubo_registrado",
    "FURTO DE VEICULO": "furto_registrado",
    "FURTO OUTROS": "furto_registrado",
    "HOMICIDIO": "homicidio_registrado",
}

CODE_RE = re.compile(r"^[A-Z]{1,2}\d{1,3}-\d+-\d+\s*-\s*")


def _normaliza(txt: str) -> str:
    txt = txt.upper()
    for de, para in (("Í", "I"), ("É", "E"), ("Ó", "O"), ("Ã", "A"), ("Â", "A"), ("Ç", "C"), ("Ô", "O")):
        txt = txt.replace(de, para)
    return txt


# ---------------------------------------------------------------------------
# 1) PRODUTIVIDADE PARAISÓPOLIS 2026 -> indicador 'produtividade_paraisopolis'
# ---------------------------------------------------------------------------
def processar_paraisopolis() -> tuple[list[dict], int, int, list[str]]:
    linhas, lidas, descartadas, descartes = [], 0, 0, []
    wb = openpyxl.load_workbook(ARQUIVO_PARAISOPOLIS, data_only=True, read_only=True)
    if "PRODUTIVIDADE" not in wb.sheetnames:
        wb.close()
        return [], 0, 0, ["aba PRODUTIVIDADE ausente em PRODUTIVIDADE PARAISÓPOLIS 2026.xlsx"]
    ws = wb["PRODUTIVIDADE"]
    soma: dict[tuple[int, int], float] = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or r[0] is None:
            continue
        data = r[0]
        lidas += 1
        if not hasattr(data, "year"):
            descartadas += 1
            descartes.append(f"produtividade_paraisopolis: DATA inválida na linha ({r[:3]})")
            continue
        pessoas = numero_ou_none(r[7]) if len(r) > 7 else None
        if pessoas is None:
            continue  # dia sem abordagem registrada, não é descarte
        chave = (data.year, data.month)
        soma[chave] = soma.get(chave, 0) + pessoas
    wb.close()
    for (ano, mes), valor in sorted(soma.items()):
        linhas.append(
            {"secao": SECAO, "indicador": "produtividade_paraisopolis", "ano": ano, "mes": mes, "eh_anual": False, "cia": 0, "valor": valor}
        )
    return linhas, lidas, descartadas, descartes


# ---------------------------------------------------------------------------
# 2) RESULTADOS OPERAÇÕES RPM 2026 -> indicador 'resultado_operacoes_rpm'
# ---------------------------------------------------------------------------
def processar_rpm() -> tuple[list[dict], int, int, list[str]]:
    linhas, lidas, descartadas, descartes = [], 0, 0, []
    wb = openpyxl.load_workbook(ARQUIVO_RPM, data_only=True, read_only=True)
    for sn in wb.sheetnames:
        sn_strip = sn.strip()
        if sn_strip in ("TOTAL", "MODELO EM BRANCO"):
            continue
        m = re.match(r"([A-Z]{3})(\d{2})", sn_strip)
        if not m:
            continue
        mes = MESES_ABREV.get(m.group(1))
        if not mes:
            continue
        ano = 2000 + int(m.group(2))
        lidas += 1
        ws = wb[sn]
        linha_pessoas = None
        for r in ws.iter_rows(values_only=True):
            if r and r[0] and str(r[0]).strip().lower() == "pessoas abordadas":
                linha_pessoas = r
                break
        if linha_pessoas is None:
            descartadas += 1
            descartes.append(f"resultado_operacoes_rpm {ano}-{mes:02d}: linha 'Pessoas abordadas' não encontrada")
            continue
        total = numero_ou_none(linha_pessoas[-1])
        if not total:  # None ou 0 -> mês sem dado real (template vazio)
            descartadas += 1
            descartes.append(f"resultado_operacoes_rpm {ano}-{mes:02d}: TOTAL vazio/zero (mês sem dado real)")
            continue
        linhas.append(
            {"secao": SECAO, "indicador": "resultado_operacoes_rpm", "ano": ano, "mes": mes, "eh_anual": False, "cia": 0, "valor": total}
        )
    wb.close()
    return linhas, lidas, descartadas, descartes


# ---------------------------------------------------------------------------
# 3) INDICADORES CRIMINAIS PARA GOOGLE MAPS -> roubo/furto/homicidio/roubo_carga
# ---------------------------------------------------------------------------
def _categoria_do_nome(nome_arquivo: str, prefixo_pasta: str) -> str:
    base = nome_arquivo[:-5] if nome_arquivo.lower().endswith(".xlsx") else nome_arquivo
    if base.startswith(prefixo_pasta):
        base = base[len(prefixo_pasta):].strip()
    base = re.sub(r"\s+(I|II|III)$", "", base).strip()
    return _normaliza(base)


def _extrai_bos(caminho: str) -> set[str]:
    """Pega a aba com mais BOs únicos (col A) — cobre tanto o layout normal
    (aba MAPS deduplicada) quanto arquivos com aba renomeada (Plan1) ou aba
    MAPS vazia por retrabalho (achado real em fev/mar26)."""
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    melhor: set[str] = set()
    for sn in wb.sheetnames:
        ws = wb[sn]
        bos = {str(r[0]).strip() for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]}
        if len(bos) > len(melhor):
            melhor = bos
    wb.close()
    return melhor


def processar_google_maps() -> tuple[list[dict], int, int, list[str], list[str]]:
    """Retorna (linhas_fato_secao, arquivos_lidos, arquivos_descartados, descartes, arquivos_processados)."""
    linhas: list[dict] = []
    descartes: list[str] = []
    arquivos_processados: list[str] = []
    arquivos_lidos = 0
    arquivos_descartados = 0

    acumulado: dict[tuple[str, int, int], set[str]] = {}

    for pasta_mes in sorted(glob.glob(os.path.join(RAIZ_GOOGLE_MAPS, "*"))):
        if not os.path.isdir(pasta_mes):
            continue
        nome_pasta = os.path.basename(pasta_mes)
        m = re.match(r"(\d{2})([A-Z]{3})(\d{2})", nome_pasta)
        if not m:
            continue
        mes_num, ano = int(m.group(1)), 2000 + int(m.group(3))
        prefixo = nome_pasta + " "
        for arq in sorted(glob.glob(os.path.join(pasta_mes, "*.xlsx"))):
            nome = os.path.basename(arq)
            if nome.startswith("~$"):
                continue
            arquivos_lidos += 1
            cat = _categoria_do_nome(nome, prefixo)
            indicador = CATEGORIA_INDICADOR.get(cat)
            if not indicador:
                arquivos_descartados += 1
                descartes.append(f"google_maps {nome_pasta}/{nome}: categoria '{cat}' não mapeada")
                continue
            try:
                bos = _extrai_bos(arq)
            except Exception as e:  # noqa: BLE001
                arquivos_descartados += 1
                descartes.append(f"google_maps {nome_pasta}/{nome}: erro ao ler ({e})")
                continue
            chave = (indicador, ano, mes_num)
            acumulado.setdefault(chave, set()).update(bos)
            arquivos_processados.append(arq)

    for (indicador, ano, mes_num), bos in sorted(acumulado.items()):
        linhas.append(
            {"secao": SECAO, "indicador": indicador, "ano": ano, "mes": mes_num, "eh_anual": False, "cia": 0, "valor": len(bos)}
        )
    return linhas, arquivos_lidos, arquivos_descartados, descartes, arquivos_processados


# ---------------------------------------------------------------------------
# 4) Snapshots Muralha (Desktop/Downloads) -> agregado_dimensional natureza
# ---------------------------------------------------------------------------
def _natureza_de(tipo: str) -> str:
    tipo = tipo.strip()
    if tipo.startswith("SSP:") or tipo.startswith("LAP:"):
        return tipo.split(":", 1)[1].strip().upper()
    sem_codigo = CODE_RE.sub("", tipo)
    base = sem_codigo.split(" - ")[0].strip().upper()
    return base or sem_codigo.strip().upper()


def processar_ocorrencias_muralha() -> tuple[list[dict], int, list[str]]:
    """Deduplica os dois snapshots (Desktop xlsx + Downloads csv, mesmo feed
    em momentos próximos) e agrega contagem por natureza."""
    eventos: dict[tuple, str] = {}
    descartes: list[str] = []

    if os.path.exists(OCORRENCIAS_XLSX):
        wb = openpyxl.load_workbook(OCORRENCIAS_XLSX, data_only=True, read_only=True)
        if "Ocorrencias" in wb.sheetnames:
            ws = wb["Ocorrencias"]
            for r in ws.iter_rows(min_row=5, values_only=True):
                if not r or not r[2]:
                    continue
                chave = (str(r[1]).strip(), str(r[2]).strip(), str(r[3]).strip())
                eventos[chave] = _natureza_de(r[2])
        else:
            descartes.append("CCO-16_Ocorrencias.xlsx: aba 'Ocorrencias' ausente")
        wb.close()
    else:
        descartes.append(f"{OCORRENCIAS_XLSX}: arquivo não encontrado")

    if os.path.exists(OCORRENCIAS_CSV):
        with open(OCORRENCIAS_CSV, encoding="utf-8-sig", errors="replace", newline="") as f:
            reader = csv.reader(f, delimiter=";")
            try:
                next(reader)  # cabeçalho
            except StopIteration:
                reader = None
            if reader is not None:
                for r in reader:
                    if len(r) < 5 or not r[2]:
                        continue
                    chave = (str(r[0]).strip(), str(r[2]).strip(), str(r[4]).strip())
                    eventos[chave] = _natureza_de(r[2])
    else:
        descartes.append(f"{OCORRENCIAS_CSV}: arquivo não encontrado")

    contagem: dict[str, int] = {}
    for natureza in eventos.values():
        contagem[natureza] = contagem.get(natureza, 0) + 1

    linhas = [
        {"secao": SECAO, "fonte": "ocorrencias_muralha", "dimensao": "natureza", "chave": natureza, "valor": valor}
        for natureza, valor in sorted(contagem.items())
    ]
    return linhas, len(eventos), descartes


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true", help="lê e mostra o resumo, não grava no Supabase")
    ap.add_argument("--forcar", action="store_true", help="reingere mesmo se o hash do arquivo não mudou")
    args = ap.parse_args()

    if not caminho_disponivel(RAIZ_PRODUTIVIDADE) or not caminho_disponivel(RAIZ_GOOGLE_MAPS):
        print(f"[p3_complementar] Fonte de rede indisponível ({RAIZ_PRODUTIVIDADE} / {RAIZ_GOOGLE_MAPS}).")
        if not args.dry_run:
            batch_id = abrir_batch(SECAO, "complementar_2026")
            fechar_batch(batch_id, "fonte_indisponivel", erro="UNC inacessível no momento da execução")
        sys.exit(0)

    batch_id = None if args.dry_run else abrir_batch(SECAO, "complementar_2026")
    total_lidas = total_validas = total_descartadas = 0
    todos_descartes: list[str] = []
    erro_fatal = None

    try:
        # --- Paraisópolis ---
        linhas_p, lidas_p, desc_p, descartes_p = processar_paraisopolis()
        total_lidas += lidas_p
        total_validas += len(linhas_p)
        total_descartadas += desc_p
        todos_descartes.extend(descartes_p)
        print(f"[p3_complementar] Paraisópolis: {lidas_p} lidas, {len(linhas_p)} meses válidos, {desc_p} descartadas")

        # --- RPM ---
        linhas_r, lidas_r, desc_r, descartes_r = processar_rpm()
        total_lidas += lidas_r
        total_validas += len(linhas_r)
        total_descartadas += desc_r
        todos_descartes.extend(descartes_r)
        print(f"[p3_complementar] RPM: {lidas_r} abas mensais lidas, {len(linhas_r)} válidas, {desc_r} descartadas")

        # --- Google Maps ---
        linhas_g, lidas_g, desc_g, descartes_g, arquivos_g = processar_google_maps()
        total_lidas += lidas_g
        total_validas += len(linhas_g)
        total_descartadas += desc_g
        todos_descartes.extend(descartes_g)
        print(f"[p3_complementar] Google Maps: {lidas_g} arquivos lidos, {len(linhas_g)} indicador-mês válidos, {desc_g} arquivos descartados")

        if not args.dry_run:
            for linha in linhas_p + linhas_r + linhas_g:
                linha["batch_id"] = batch_id
            upsert_fato_secao(linhas_p + linhas_r + linhas_g)

            for caminho in [ARQUIVO_PARAISOPOLIS, ARQUIVO_RPM] + arquivos_g:
                sha = sha256_arquivo(caminho)
                registrar_arquivo(
                    secao=SECAO,
                    caminho_unc=caminho,
                    sha256=sha,
                    mtime_iso=mtime_iso(caminho),
                    linhas_reais=None,
                    observacao="P3 complementar 2026 — produtividade Paraisópolis / RPM / indicadores criminais Google Maps",
                    batch_id=batch_id,
                )

        # --- Ocorrências Muralha (snapshot local) ---
        linhas_m, n_eventos_m, descartes_m = processar_ocorrencias_muralha()
        total_lidas += n_eventos_m
        total_validas += len(linhas_m)
        todos_descartes.extend(descartes_m)
        print(f"[p3_complementar] Ocorrências Muralha: {n_eventos_m} eventos deduplicados, {len(linhas_m)} naturezas")

        if not args.dry_run and linhas_m:
            upsert_agregado_dimensional(linhas_m)
            for caminho in (OCORRENCIAS_XLSX, OCORRENCIAS_CSV):
                if os.path.exists(caminho):
                    sha = sha256_arquivo(caminho)
                    registrar_arquivo(
                        secao=SECAO,
                        caminho_unc=caminho,
                        sha256=sha,
                        mtime_iso=mtime_iso(caminho),
                        linhas_reais=n_eventos_m,
                        observacao="Snapshot local Muralha Paulista (janela móvel 3 dias) — agregado por natureza",
                        batch_id=batch_id,
                    )

    except Exception as e:  # noqa: BLE001 — precisa fechar o batch como 'falha' em qualquer exceção
        erro_fatal = str(e)

    if args.dry_run:
        print(f"[p3_complementar] dry-run concluído: {total_lidas} lidas, {total_validas} válidas, {total_descartadas} descartadas.")
        if todos_descartes:
            print("  Amostra de descartes:", todos_descartes[:10])
        return

    if erro_fatal:
        fechar_batch(batch_id, "falha", total_lidas, total_validas, total_descartadas, todos_descartes, erro_fatal)
        print(f"[p3_complementar] FALHA: {erro_fatal}")
        sys.exit(1)

    status = "ok" if total_descartadas == 0 else "parcial"
    fechar_batch(batch_id, status, total_lidas, total_validas, total_descartadas, todos_descartes)
    print(f"[p3_complementar] Concluído ({status}): {total_lidas} lidas, {total_validas} válidas, {total_descartadas} descartadas.")


if __name__ == "__main__":
    main()
