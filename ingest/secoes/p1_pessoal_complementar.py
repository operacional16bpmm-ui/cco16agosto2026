"""
Ingestão P1 — complementar ao QSE (LTS + PAF/férias).

NÃO mexe nos indicadores QSE já ingeridos por p1_qse.py (efetivo_existente,
efetivo_fixado, efetivo_apto, efetivo_restricao_medica,
efetivo_restricao_operacional, efetivo_inapto) — usa chaves de indicador
novas e distintas: 'lts_afastamentos' e 'ferias_concedidas'.

Fonte 1 — LTS (Licença para Tratamento de Saúde), relatório oficial da
Diretoria de Pessoal (SSP-SP), já filtrado por OPM:
    \\cmdo\pmesp\16BPMM\16BPMM_EM\P1\Matriz\Matriz 2025\
    INDICADORES PLANEJ ESTRATÉGICO - DRIVE CPA\LTS <ano ou mês>.xls
    (aba "restricaoRelatorio", formato .xls legado — xlrd)
Cobertura real encontrada: 2022, 2023 e 2024 completos (arquivos "LTS
<ano>.xls"); 2025 só até março (arquivos "LTS JAN/FEV/MAR.xls" — não existe
"LTS 2025.xls" nem dados de abr-dez/2025 nesta pasta). Cada linha de dado
tem 8 células não-vazias na ordem Posto/Grad, RE, Nome, OPM, Tipo Restrição,
DT Início, DT Término, Qtde Dias (a coluna "Histórico" do cabeçalho vem
sempre vazia nos dados, por isso não aparece nas 8 células). Filtra por
OPM contendo "16" + "BPM" (relatório já vem só com 16.BPM/M nas amostras
conferidas, mas o filtro fica explícito por segurança). Agrega por ano de
DT Início → indicador 'lts_afastamentos', eh_anual=true, cia=0.

Fonte 2 — PAF (Planejamento Anual de Férias), consolidado do batalhão
(todas as seções numa aba só, não as abas por seção que são subconjunto):
    Matriz 2025\PAF 2025\PAF 2025.xlsx        (aba "2024" — rótulo da aba
        não é confiável: as datas de dentro são todas de 2025, achado real
        na conferência linha a linha)
    Matriz 2025\PAF 2026\PAF_2026 (1).xlsx    (aba "2026" — datas de 2026
        e um resíduo de dez/2025)
Colunas fixas nas duas: [0]=GRADUAÇÃO [1]=RE [2]=NOME [8]=1ª OPÇÃO (data
início) [9]=2ª OPÇÃO (data fim, pode faltar) [10]=dias (int ou texto "NN
DIAS"). Não achei PAF consolidado de batalhão para 2024 (só um rascunho com
14 linhas válidas e uma planilha isolada da seção P4 em "FÉRIAS 2023") —
registrado em nao_carregado pelo caller. Dedupe global por (RE, data
início) antes de agregar, porque a mesma concessão de férias aparece em
mais de um arquivo/revisão (ex.: 12 registros de 2025 repetidos dentro do
arquivo de 2026). Agrega por (ano, mês) da data de início → indicador
'ferias_concedidas', eh_anual=false, cia=0.

Uso:
    python -m ingest.secoes.p1_pessoal_complementar [--dry-run] [--forcar]
"""
import argparse
import datetime as dt
import os
import re
import sys

import openpyxl
import xlrd

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
from ingest.common.unc import caminho_disponivel
from ingest.common.excel import mtime_iso
from ingest.common.carga import (
    abrir_batch,
    fechar_batch,
    upsert_fato_secao,
    registrar_arquivo,
    arquivo_ja_ingerido,
    sha256_arquivo,
)

SECAO = "p1"
RAIZ_UNC = r"\\cmdo\pmesp\16BPMM\16BPMM_EM\P1\Matriz"

LTS_ARQUIVOS = [
    r"Matriz 2025\INDICADORES PLANEJ ESTRATÉGICO - DRIVE CPA\LTS 2022.xls",
    r"Matriz 2025\INDICADORES PLANEJ ESTRATÉGICO - DRIVE CPA\LTS 2023.xls",
    r"Matriz 2025\INDICADORES PLANEJ ESTRATÉGICO - DRIVE CPA\LTS 2024.xls",
    r"Matriz 2025\INDICADORES PLANEJ ESTRATÉGICO - DRIVE CPA\LTS JAN.xls",
    r"Matriz 2025\INDICADORES PLANEJ ESTRATÉGICO - DRIVE CPA\LTS FEV.xls",
    r"Matriz 2025\INDICADORES PLANEJ ESTRATÉGICO - DRIVE CPA\LTS MAR.xls",
]

# (caminho relativo, nome da aba) — só os consolidados de batalhão, não as
# abas por seção (P1/P2/P3/P4/P5/MOTOMEC/SPJMD) que são subconjunto delas.
PAF_ARQUIVOS = [
    (r"Matriz 2025\PAF 2025\PAF 2025.xlsx", "2024"),
    (r"Matriz 2025\PAF 2026\PAF_2026 (1).xlsx", "2026"),
]

RE_DATA_BR = re.compile(r"^\d{2}/\d{2}/\d{4}$")


# ---------------------------------------------------------------- LTS ----

def extrair_lts(caminho: str) -> list[dict]:
    """Retorna registros brutos [{re, dt_inicio, dt_termino, dias}] de UMA
    aba restricaoRelatorio. dt_inicio/dt_termino ficam como string dd/mm/aaaa
    (não convertidas ainda) para servir de chave de dedupe estável."""
    wb = xlrd.open_workbook(caminho)
    if "restricaoRelatorio" not in wb.sheet_names():
        return []
    ws = wb.sheet_by_name("restricaoRelatorio")
    registros = []
    for r in range(ws.nrows):
        celulas = [str(ws.cell_value(r, c)).strip() for c in range(ws.ncols)]
        celulas = [v for v in celulas if v]
        if len(celulas) != 8:
            continue
        opm, dt_ini, dt_fim, dias = celulas[3], celulas[5], celulas[6], celulas[7]
        if "BPM" not in opm.upper() or "16" not in opm:
            continue
        if not (RE_DATA_BR.match(dt_ini) and RE_DATA_BR.match(dt_fim) and dias.isdigit()):
            continue
        registros.append({"re": celulas[1], "dt_inicio": dt_ini, "dt_termino": dt_fim, "dias": int(dias)})
    return registros


def processar_lts(dry_run: bool) -> tuple[list[dict], int, int, list[str], list[tuple[str, int]]]:
    """Retorna (linhas_fato_secao, lidas, descartadas, descartes,
    [(caminho_abs, n_registros_brutos)])."""
    vistos: dict[tuple, dict] = {}
    lidas = 0
    descartes: list[str] = []
    arquivos_processados: list[tuple[str, int]] = []

    for rel in LTS_ARQUIVOS:
        caminho = os.path.join(RAIZ_UNC, rel)
        if not os.path.isfile(caminho):
            descartes.append(f"LTS: arquivo não encontrado — {caminho}")
            continue
        try:
            registros = extrair_lts(caminho)
        except Exception as e:  # noqa: BLE001
            descartes.append(f"LTS: falha ao ler {os.path.basename(caminho)}: {e}")
            continue
        arquivos_processados.append((caminho, len(registros)))
        lidas += len(registros)
        for reg in registros:
            chave = (reg["re"], reg["dt_inicio"], reg["dt_termino"])
            if chave in vistos:
                continue
            vistos[chave] = reg

    descartadas = lidas - len(vistos)
    if descartadas:
        descartes.append(f"LTS: {descartadas} registro(s) duplicado(s) entre arquivos (mesmo RE+período) — deduplicados")

    contagem: dict[int, int] = {}
    for reg in vistos.values():
        ano = int(reg["dt_inicio"][-4:])
        contagem[ano] = contagem.get(ano, 0) + 1

    linhas = [
        {
            "secao": SECAO,
            "indicador": "lts_afastamentos",
            "ano": ano,
            "mes": None,
            "eh_anual": True,
            "cia": 0,
            "valor": total,
        }
        for ano, total in sorted(contagem.items())
    ]
    return linhas, lidas, descartadas, descartes, arquivos_processados


# ---------------------------------------------------------------- PAF ----

def _dias_de(valor) -> int | None:
    if isinstance(valor, (int, float)):
        return int(valor)
    if isinstance(valor, str):
        m = re.search(r"\d+", valor)
        if m:
            return int(m.group())
    return None


def extrair_paf(caminho: str, aba: str) -> list[dict]:
    """Retorna registros brutos [{re, dt_inicio(date), dt_termino, dias}]."""
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    try:
        if aba not in wb.sheetnames:
            return []
        ws = wb[aba]
        registros = []
        for row in ws.iter_rows(values_only=True):
            if len(row) < 11:
                continue
            re_bruto = row[1]
            data1 = row[8]
            if not isinstance(re_bruto, str) or not re.match(r"^\d", re_bruto.strip()):
                continue
            if not isinstance(data1, dt.datetime):
                continue
            data2 = row[9] if isinstance(row[9], dt.datetime) else None
            registros.append(
                {
                    "re": re_bruto.strip(),
                    "dt_inicio": data1.date(),
                    "dt_termino": data2.date() if data2 else None,
                    "dias": _dias_de(row[10]),
                }
            )
        return registros
    finally:
        wb.close()


def processar_paf(dry_run: bool) -> tuple[list[dict], int, int, list[str], list[tuple[str, int]]]:
    vistos: dict[tuple, dict] = {}
    lidas = 0
    descartes: list[str] = []
    arquivos_processados: list[tuple[str, int]] = []

    for rel, aba in PAF_ARQUIVOS:
        caminho = os.path.join(RAIZ_UNC, rel)
        if not os.path.isfile(caminho):
            descartes.append(f"PAF: arquivo não encontrado — {caminho}")
            continue
        try:
            registros = extrair_paf(caminho, aba)
        except Exception as e:  # noqa: BLE001
            descartes.append(f"PAF: falha ao ler {os.path.basename(caminho)} aba {aba}: {e}")
            continue
        if not registros:
            descartes.append(f"PAF: aba '{aba}' ausente ou vazia em {os.path.basename(caminho)}")
            continue
        arquivos_processados.append((caminho, len(registros)))
        lidas += len(registros)
        for reg in registros:
            chave = (reg["re"], reg["dt_inicio"])
            if chave in vistos:
                continue
            vistos[chave] = reg

    descartadas = lidas - len(vistos)
    if descartadas:
        descartes.append(f"PAF: {descartadas} registro(s) duplicado(s) entre arquivos (mesmo RE+data início) — deduplicados")

    contagem: dict[tuple[int, int], int] = {}
    for reg in vistos.values():
        chave = (reg["dt_inicio"].year, reg["dt_inicio"].month)
        contagem[chave] = contagem.get(chave, 0) + 1

    linhas = [
        {
            "secao": SECAO,
            "indicador": "ferias_concedidas",
            "ano": ano,
            "mes": mes,
            "eh_anual": False,
            "cia": 0,
            "valor": total,
        }
        for (ano, mes), total in sorted(contagem.items())
    ]
    return linhas, lidas, descartadas, descartes, arquivos_processados


# --------------------------------------------------------------- main ----

def _rodar_fonte(nome_fonte: str, processar_fn, observacao: str, args) -> tuple[int, int, int, list[str]]:
    linhas, lidas, descartadas, descartes, arquivos = processar_fn(args.dry_run)
    print(f"[p1_pessoal_complementar] fonte={nome_fonte}: {lidas} lidas, {len(linhas)} agregados, {descartadas} descartadas")
    for d in descartes[:10]:
        print(f"  - {d}")

    if args.dry_run:
        return lidas, len(linhas), descartadas, descartes

    if not linhas and not arquivos:
        batch_id = abrir_batch(SECAO, nome_fonte)
        fechar_batch(batch_id, "falha", erro="nenhum arquivo/registro localizado")
        return lidas, len(linhas), descartadas, descartes

    batch_id = abrir_batch(SECAO, nome_fonte)
    try:
        for linha in linhas:
            linha["batch_id"] = batch_id
        upsert_fato_secao(linhas)
        for caminho, n_registros in arquivos:
            sha = sha256_arquivo(caminho)
            registrar_arquivo(
                secao=SECAO,
                caminho_unc=caminho,
                sha256=sha,
                mtime_iso=mtime_iso(caminho),
                linhas_reais=n_registros,
                observacao=observacao,
                batch_id=batch_id,
            )
    except Exception as e:  # noqa: BLE001
        fechar_batch(batch_id, "falha", lidas, len(linhas), descartadas, descartes, erro=str(e))
        raise

    status = "ok" if descartadas == 0 else "parcial"
    fechar_batch(batch_id, status, lidas, len(linhas), descartadas, descartes)
    return lidas, len(linhas), descartadas, descartes


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--forcar", action="store_true", help="mantido por consistência de CLI; sem efeito aqui (arquivos fixos, sempre reprocessados)")
    args = ap.parse_args()

    if not caminho_disponivel(RAIZ_UNC):
        print(f"[p1_pessoal_complementar] Fonte indisponível: {RAIZ_UNC}")
        if not args.dry_run:
            for fonte in ("lts_afastamentos", "paf_ferias"):
                batch_id = abrir_batch(SECAO, fonte)
                fechar_batch(batch_id, "fonte_indisponivel", erro="UNC inacessível no momento da execução")
        sys.exit(0)

    total_lidas = total_validas = total_descartadas = 0

    l_lidas, l_validas, l_desc, _ = _rodar_fonte(
        "lts_afastamentos", processar_lts,
        "LTS — Licença Tratamento Saúde, contagem anual de afastamentos (relatório DP/SSP-SP)", args,
    )
    p_lidas, p_validas, p_desc, _ = _rodar_fonte(
        "paf_ferias", processar_paf,
        "PAF — Planejamento Anual de Férias, concessões mensais consolidadas do batalhão", args,
    )

    total_lidas = l_lidas + p_lidas
    total_validas = l_validas + p_validas
    total_descartadas = l_desc + p_desc

    if args.dry_run:
        print(f"[p1_pessoal_complementar] dry-run concluído: {total_lidas} lidas, {total_validas} agregados válidos, {total_descartadas} descartadas.")
        return

    print(f"[p1_pessoal_complementar] Concluído: {total_lidas} lidas, {total_validas} agregados válidos, {total_descartadas} descartadas.")


if __name__ == "__main__":
    main()
