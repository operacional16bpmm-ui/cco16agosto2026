"""
Ingestão P5 — Reconhecimento/Mérito (LMP, Agraciados, Indicados, Campanhas
Institucionais) + indicador de governança (Objetivos Estratégicos).

Fonte: \\cmdo\pmesp\16BPMM\16BPMM_EM\P5\...
  - LMP mensal agregado:
      2026 - Jéssica e Ramaldes\Láurea do Mérito Pessoal - LMP\
        QUANTITATIVO MENSAL DE LMP PAGAS.xlsx (aba "Controle de Láureas")
  - LMP nominal trimestral:
      2026 - Jéssica e Ramaldes\Láurea do Mérito Pessoal - LMP\
        Controle trimestral de LMP - JAN FEV MAR.xlsx (aba "JUN26")
      (o arquivo irmão em LAUREAS SGT MIRIAM\Controle trimestral de LMP -
       ABR MAI JUN.xlsx é cópia idêntica — mesmas 3 seções de mês/ano — só o
       primeiro é lido, para não duplicar nominal.)
  - Agraciados Challenge Coin (2023-2026, acumulado):
      2026\Challenge Coin\AGRACIADOS 2023 à 2026.xlsx (aba "2023 a 2026")
  - Agraciados Medalha do Cinquentenário (2013-2026, todas as seções "X anos"):
      2025\CINQUENTENÁRIO\MEDALHA CONCEDIDA EM TODOS ANOS.xlsx (aba "Plan1")
      Achado real: esta planilha tem ~1.048.494 linhas de formatação vazia
      após a última seção real (63 anos/2026) — a leitura para no primeiro
      bloco de 40 linhas vazias consecutivas, não varre o arquivo inteiro.
  - Indicados (Medalha Cinquentenário 2025):
      2025\CINQUENTENÁRIO\INDICADOS 2025\GERAL - INDICADOS OFICIAL.xlsx
        (aba "por OPM" — a aba "por antiguidade" é a MESMA lista, só
        reordenada por outro critério; carregar as duas duplicaria cada
        pessoa, então só "por OPM" é lida)
      2025\CINQUENTENÁRIO\INDICADOS 2025\banco_dados_INDICADOS.xlsx
        (aba "INDICADOS 2025" — banco de candidatos, universo maior que a
        lista oficial acima; carregado com objetivo/status distintos para
        não colidir semanticamente com os indicados oficiais)
  - Campanhas institucionais (2022-2026): dezenas de arquivos mensais
      "*CAMPANHAS INSTITUCIONAIS*.xlsx" / "Quadro Campanhas Institucionais
      *.xlsx" espalhados em "Documentos de Prazo\...\(dia 20/30) Campanhas
      Institucionais\<mês>\". Cada arquivo tem uma aba por área CPAM
      (CPAM-1..CPAM-11, CPC) listando OPMs da área; só a aba que contém a
      linha "16º BPM/M" preenchida (varia entre CPAM-1 e CPAM-5 conforme o
      arquivo) tem dado real do batalhão — as outras ~10 abas por arquivo
      são de outras unidades e são ignoradas. O mês/ano real está na célula
      "MÊS DE REFERÊNCIA - <mês>[/<ano>]" da própria aba, não no nome do
      arquivo/pasta (achado real: há arquivos arquivados na pasta "Janeiro"
      cujo conteúdo interno é "MÊS DE REFERÊNCIA - JUNHO" — pasta indica só
      quando foi arquivado, não o mês do dado). Dezenas dessas planilhas são
      cópias redundantes do mesmo mês (mesmo relatório reenviado por
      Cia/CPAM) — deduplicado por (ano, mês, nome da campanha), mantendo a
      maior quantidade observada entre as cópias.
  - Indicadores dos Objetivos Estratégicos (GOVERNANÇA, não reconhecimento):
      2025\Documento de Prazo\TRIMESTREAL JUN Cópia de Indicadores dos
        Objetivos Estratégicos JUN.xlsx (aba "C Com Soc") — usada em vez da
        irmã "TRIMESTREAL - Indicadores dos Objetivos Estratégicos.xlsx"
        porque tem dado até outubro/2025 (a outra só vai até junho) e é a
        cópia mais completa da mesma tabela.
      Grava em agregado_dimensional (secao='governanca', sem restrição de
      schema). NÃO grava em fato_secao: a coluna fato_secao.secao tem CHECK
      restrito a ('p1','p2','p3','p4','p5','ft','motomec','res_armas',
      'spjmd') — 'governanca' violaria a constraint. Alterar essa constraint
      exigiria uma migration nova, e por instrução do projeto (memória
      "Autorização prévia") nenhum arquivo de schema é criado/alterado sem
      perguntar antes — então esta parte fica pendente de autorização
      (reportada em nao_carregado/erros) em vez de ingerida.

Tabelas dedicadas (migration 011) sem unique constraint além de id — este
script só faz POST (nunca upsert) e só depois de conferir com SELECT count
que a tabela-alvo está vazia, para não duplicar em reruns.

Uso:
    python -m ingest.secoes.p5_reconhecimento [--dry-run]
"""
import argparse
import glob
import os
import re
import sys
import unicodedata

import openpyxl
import requests

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
from ingest.common.unc import caminho_disponivel
from ingest.common.excel import eh_arquivo_lixo, mtime_iso
from ingest.common.carga import (
    abrir_batch,
    fechar_batch,
    upsert_generico,
    upsert_agregado_dimensional,
    sha256_arquivo,
    _HEADERS,
    _url,
)

SECAO = "p5"
FONTE = "p5_reconhecimento"
RAIZ_UNC = r"\\cmdo\pmesp\16BPMM\16BPMM_EM\P5"

MESES_PT = [
    "JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO",
    "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO",
]
MESES_ABV = ["JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT", "NOV", "DEZ"]

RE_RE = re.compile(r"^\d{4,9}-?[0-9A-Za-z]?$")


def _norm(s) -> str:
    s = "" if s is None else str(s)
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    return s.upper().strip()


def _celula(v) -> str:
    return "" if v is None else str(v).strip()


def normalizar_chaves(linhas: list[dict]) -> list[dict]:
    """PostgREST (bulk insert) exige que todo objeto do array JSON tenha
    exatamente o mesmo conjunto de chaves — dicts com chaves variáveis (ex.:
    linha agregada mensal sem 're'/'cia'/'extra' vs. linha nominal com essas
    chaves) disparam erro PGRST102 'All object keys must match'. Preenche
    com None as chaves ausentes em cada linha, usando a união de todas."""
    if not linhas:
        return linhas
    todas_chaves = set()
    for l in linhas:
        todas_chaves.update(l.keys())
    return [{k: l.get(k) for k in todas_chaves} for l in linhas]


def tabela_vazia(tabela: str) -> bool:
    resp = requests.get(
        _url(tabela),
        headers=_HEADERS,
        params={"select": "id", "limit": "1"},
        timeout=30,
    )
    resp.raise_for_status()
    return len(resp.json()) == 0


def registrar_arquivo_p5(dataset: str, caminho_unc: str, aba: str, linhas_reais: int, observacao: str = None):
    upsert_generico(
        "p5_arquivos_fonte",
        [{
            "dataset": dataset,
            "nome_arquivo": os.path.basename(caminho_unc),
            "caminho_unc": caminho_unc,
            "aba": aba,
            "linhas_reais": linhas_reais,
            "observacao": observacao,
        }],
        on_conflict="dataset,caminho_unc,aba",
    )


# ----------------------------------------------------------------------------
# LMP
# ----------------------------------------------------------------------------

def carregar_lmp(descartes: list) -> list[dict]:
    linhas = []

    # -- (a) agregado mensal: QUANTITATIVO MENSAL DE LMP PAGAS.xlsx --------
    caminho = os.path.join(
        RAIZ_UNC, "2026 - Jéssica e Ramaldes", "Láurea do Mérito Pessoal - LMP",
        "QUANTITATIVO MENSAL DE LMP PAGAS.xlsx",
    )
    if os.path.isfile(caminho):
        wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
        aba = "Controle de Láureas"
        if aba in wb.sheetnames:
            ws = wb[aba]
            rows = list(ws.iter_rows(values_only=True))
            colunas_grau = {1: 1, 2: 2, 3: 3, 4: 4, 5: 5}  # índice coluna -> grau
            for r in rows[1:]:  # pula cabeçalho
                mes_txt = _norm(r[0] if len(r) > 0 else None)
                if mes_txt not in MESES_PT:
                    continue
                mes_num = MESES_PT.index(mes_txt) + 1
                total = r[7] if len(r) > 7 else None
                if not isinstance(total, (int, float)) or total <= 0:
                    continue  # mês sem láurea paga/publicada — nada a gravar
                for grau, idx in colunas_grau.items():
                    v = r[idx] if len(r) > idx else None
                    if isinstance(v, (int, float)) and v > 0:
                        linhas.append({
                            "ano": 2026, "mes": mes_num, "trimestre": None,
                            "tipo_lmp": f"{grau}º grau", "status": "paga/publicada",
                            "quantidade": int(v),
                            "origem_arquivo": os.path.basename(caminho), "aba": aba,
                        })
                cc = r[6] if len(r) > 6 else None
                if isinstance(cc, (int, float)) and cc > 0:
                    linhas.append({
                        "ano": 2026, "mes": mes_num, "trimestre": None,
                        "tipo_lmp": "challenge coin", "status": "paga/publicada",
                        "quantidade": int(cc),
                        "origem_arquivo": os.path.basename(caminho), "aba": aba,
                    })
            registrar_arquivo_p5("lmp", caminho, aba, len(rows), "agregado mensal por grau, 2026")
        else:
            descartes.append(f"LMP quantitativo: aba '{aba}' ausente em {caminho}")
        wb.close()
    else:
        descartes.append(f"LMP quantitativo: arquivo não encontrado {caminho}")

    # -- (b) nominal trimestral: Controle trimestral de LMP - JAN FEV MAR --
    caminho = os.path.join(
        RAIZ_UNC, "2026 - Jéssica e Ramaldes", "Láurea do Mérito Pessoal - LMP",
        "Controle trimestral de LMP - JAN FEV MAR.xlsx",
    )
    if os.path.isfile(caminho):
        wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
        aba = "JUN26"
        if aba in wb.sheetnames:
            ws = wb[aba]
            rows = list(ws.iter_rows(values_only=True))
            ano_atual, mes_atual = None, None
            lidas = 0
            for r in rows:
                c0 = r[0] if len(r) > 0 else None
                if hasattr(c0, "year"):  # célula de data = cabeçalho de seção do mês
                    ano_atual, mes_atual = c0.year, c0.month
                    continue
                nome = _celula(c0)
                if not nome or mes_atual is None or _norm(nome) == "NOME":
                    continue  # "NOME" = linha de cabeçalho de coluna repetida a cada seção de mês
                lidas += 1
                posto = _celula(r[1] if len(r) > 1 else None)
                grau = r[2] if len(r) > 2 else None
                boletim = _celula(r[3] if len(r) > 3 else None)
                data_bol = r[4] if len(r) > 4 else None
                m_re = re.search(r"\b(\d{5,6}-?[0-9A-Za-z])\b", nome)
                re_val = m_re.group(1) if m_re else None
                nome_limpo = nome.replace(m_re.group(1), "").strip() if m_re else nome
                nome_limpo = re.sub(r"^(SD|CB|CABO|SGT|1º|2º|3º|TEN|CAP|MAJ)\.?\s*PM\.?\s*", "", nome_limpo, flags=re.IGNORECASE).strip()
                if not nome_limpo:
                    nome_limpo = nome
                try:
                    grau_num = int(grau) if grau is not None else None
                except (TypeError, ValueError):
                    grau_num = None
                linhas.append({
                    "ano": ano_atual, "mes": mes_atual, "trimestre": None,
                    "re": re_val, "nome_guerra": nome_limpo, "cia": posto or None,
                    "tipo_lmp": f"{grau_num}º grau" if grau_num else None,
                    "status": "publicada" if boletim else None,
                    "quantidade": None,
                    "origem_arquivo": os.path.basename(caminho), "aba": aba,
                    "extra": {"boletim": boletim or None, "data_boletim": str(data_bol) if data_bol else None},
                })
            registrar_arquivo_p5("lmp", caminho, aba, lidas, "nominal, seções abr/mai/jun 2026")
        else:
            descartes.append(f"LMP trimestral: aba '{aba}' ausente em {caminho}")
        wb.close()
    else:
        descartes.append(f"LMP trimestral: arquivo não encontrado {caminho}")

    descartes.append(
        "LMP trimestral: LAUREAS SGT MIRIAM\\Controle trimestral de LMP - ABR MAI JUN.xlsx "
        "é cópia idêntica do arquivo já lido (mesmas 3 seções jun/abr/mai 2026) — não lido de novo p/ não duplicar."
    )
    return linhas


# ----------------------------------------------------------------------------
# AGRACIADOS
# ----------------------------------------------------------------------------

def carregar_agraciados(descartes: list) -> list[dict]:
    linhas = []

    # -- (a) Challenge Coin 2023-2026 ---------------------------------------
    caminho = os.path.join(RAIZ_UNC, "2026", "Challenge Coin", "AGRACIADOS 2023 à 2026.xlsx")
    if os.path.isfile(caminho):
        wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
        aba = "2023 a 2026"
        if aba in wb.sheetnames:
            ws = wb[aba]
            rows = list(ws.iter_rows(values_only=True))
            ano_atual = None
            lidas = 0
            for r in rows:
                col_txt = " ".join(_celula(c) for c in r)
                if "BOLETIM" in _norm(col_txt):
                    m = re.search(r"(\d{2})[A-ZÇ]{3}(\d{2})", _norm(col_txt))
                    if m:
                        ano_atual = 2000 + int(m.group(2))
                    else:
                        m2 = re.search(r"/(\d{2})\b", col_txt)
                        if m2:
                            ano_atual = 2000 + int(m2.group(1))
                    continue
                posto = _celula(r[1] if len(r) > 1 else None)
                re_val = _celula(r[2] if len(r) > 2 else None)
                nome = _celula(r[3] if len(r) > 3 else None)
                unidade = _celula(r[4] if len(r) > 4 else None)
                if not RE_RE.match(re_val) or not nome or not posto:
                    continue
                lidas += 1
                linhas.append({
                    "ano": ano_atual, "re": re_val, "nome_guerra": nome, "cia": unidade or None,
                    "medalha": "Challenge Coin 16º BPM/M", "tipo": "Challenge Coin",
                    "data_concessao": None,
                    "origem_arquivo": os.path.basename(caminho), "aba": aba,
                    "extra": {"posto_grad": posto},
                })
            registrar_arquivo_p5("agraciados", caminho, aba, lidas, "Challenge Coin, acumulado 2023-2026")
        else:
            descartes.append(f"Agraciados Challenge Coin: aba '{aba}' ausente em {caminho}")
        wb.close()
    else:
        descartes.append(f"Agraciados Challenge Coin: arquivo não encontrado {caminho}")

    # -- (b) Medalha do Cinquentenário (todas as seções "X anos") ----------
    caminho = os.path.join(RAIZ_UNC, "2025", "CINQUENTENÁRIO", "MEDALHA CONCEDIDA EM TODOS ANOS.xlsx")
    if os.path.isfile(caminho):
        wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
        aba = "Plan1"
        if aba in wb.sheetnames:
            ws = wb[aba]
            ano_atual, tipo_atual = None, None
            lidas = 0
            linhas_vazias_seguidas = 0
            LIMITE_VAZIAS = 40  # após a última seção real (63 anos/2026) só há
            # ~1M linhas de formatação vazia (achado real, ver docstring) — para
            # de ler ao encontrar essa quantidade de linhas em branco seguidas.
            for r in ws.iter_rows(max_col=6, values_only=True):
                vazio = not any(_celula(c) for c in r)
                if vazio:
                    linhas_vazias_seguidas += 1
                    if linhas_vazias_seguidas >= LIMITE_VAZIAS:
                        break
                    continue
                linhas_vazias_seguidas = 0
                col1 = _norm(r[1] if len(r) > 1 else None)
                if re.match(r"^\d+\s*ANOS$", col1):
                    ano_txt = r[2] if len(r) > 2 else None
                    try:
                        ano_atual = int(ano_txt)
                    except (TypeError, ValueError):
                        pass
                    tipo_atual = _celula(r[1])
                    continue
                posto = _celula(r[1] if len(r) > 1 else None)
                re_val = _celula(r[2] if len(r) > 2 else None)
                nome = _celula(r[3] if len(r) > 3 else None)
                opm = _celula(r[4] if len(r) > 4 else None)
                funcao = _celula(r[5] if len(r) > 5 else None)
                if not RE_RE.match(re_val) or not nome or not posto:
                    continue
                lidas += 1
                linhas.append({
                    "ano": ano_atual, "re": re_val, "nome_guerra": nome, "cia": opm or None,
                    "medalha": "Medalha do Cinquentenário do 16º BPM/M", "tipo": tipo_atual,
                    "data_concessao": None,
                    "origem_arquivo": os.path.basename(caminho), "aba": aba,
                    "extra": {"posto_grad": posto, "funcao_exercida": funcao or None},
                })
            registrar_arquivo_p5(
                "agraciados", caminho, aba, lidas,
                f"Medalha Cinquentenário, seções 2013-2026; leitura truncada após "
                f"{LIMITE_VAZIAS} linhas vazias seguidas (aba-fantasma de ~1M linhas de formatação)",
            )
        else:
            descartes.append(f"Agraciados Medalha Cinquentenário: aba '{aba}' ausente em {caminho}")
        wb.close()
    else:
        descartes.append(f"Agraciados Medalha Cinquentenário: arquivo não encontrado {caminho}")

    return linhas


# ----------------------------------------------------------------------------
# INDICADOS
# ----------------------------------------------------------------------------

def carregar_indicados(descartes: list) -> list[dict]:
    linhas = []

    # -- (a) GERAL - INDICADOS OFICIAL.xlsx, aba "por OPM" ------------------
    caminho = os.path.join(RAIZ_UNC, "2025", "CINQUENTENÁRIO", "INDICADOS 2025", "GERAL - INDICADOS OFICIAL.xlsx")
    if os.path.isfile(caminho):
        wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
        aba = "por OPM"
        if aba in wb.sheetnames:
            ws = wb[aba]
            rows = list(ws.iter_rows(values_only=True))
            cia_atual = None
            lidas = 0
            for r in rows[6:]:  # cabeçalho na linha 6 (0-indexed 5)
                cia_cel = _celula(r[0] if len(r) > 0 else None)
                if cia_cel:
                    cia_atual = cia_cel
                posto = _celula(r[1] if len(r) > 1 else None)
                re_val = _celula(r[2] if len(r) > 2 else None)
                nome = _celula(r[3] if len(r) > 3 else None)
                if not RE_RE.match(re_val) or not nome:
                    continue
                lidas += 1
                comportamento = _celula(r[5] if len(r) > 5 else None)
                opm = _celula(r[7] if len(r) > 7 else None)
                linhas.append({
                    "ano": 2025, "re": re_val, "nome_guerra": nome, "cia": cia_atual,
                    "objetivo": "Medalha Cinquentenário 16º BPM/M",
                    "status": "indicado_oficial",
                    "origem_arquivo": os.path.basename(caminho), "aba": aba,
                    "extra": {
                        "posto_grad": posto, "comportamento": comportamento or None,
                        "opm": opm or None,
                        "tempo_16m": r[4] if len(r) > 4 else None,
                    },
                })
            registrar_arquivo_p5("indicados", caminho, aba, lidas, "indicados oficiais Medalha Cinquentenário 2025, por OPM/Cia")
        else:
            descartes.append(f"Indicados oficiais: aba '{aba}' ausente em {caminho}")
        wb.close()
        descartes.append(
            "Indicados oficiais: aba 'por antiguidade' do mesmo arquivo é a MESMA lista "
            "(mesmas 59 pessoas), só reordenada — não lida de novo p/ não duplicar."
        )
    else:
        descartes.append(f"Indicados oficiais: arquivo não encontrado {caminho}")

    # -- (b) banco_dados_INDICADOS.xlsx --------------------------------------
    caminho = os.path.join(RAIZ_UNC, "2025", "CINQUENTENÁRIO", "INDICADOS 2025", "banco_dados_INDICADOS.xlsx")
    if os.path.isfile(caminho):
        wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
        aba = "INDICADOS 2025"
        if aba in wb.sheetnames:
            ws = wb[aba]
            rows = list(ws.iter_rows(values_only=True))
            lidas = 0
            for r in rows[1:]:
                posto = _celula(r[1] if len(r) > 1 else None)
                re_val = _celula(r[2] if len(r) > 2 else None)
                nome = _celula(r[3] if len(r) > 3 else None)
                if not RE_RE.match(re_val) or not nome:
                    continue
                lidas += 1
                linhas.append({
                    "ano": 2025, "re": re_val, "nome_guerra": nome, "cia": None,
                    "objetivo": "Medalha Cinquentenário 16º BPM/M — banco de dados",
                    "status": "banco_dados",
                    "origem_arquivo": os.path.basename(caminho), "aba": aba,
                    "extra": {"posto_grad": posto},
                })
            registrar_arquivo_p5("indicados", caminho, aba, lidas, "banco de candidatos Medalha Cinquentenário 2025 (universo maior que a lista oficial)")
        else:
            descartes.append(f"Banco de dados indicados: aba '{aba}' ausente em {caminho}")
        wb.close()
    else:
        descartes.append(f"Banco de dados indicados: arquivo não encontrado {caminho}")

    return linhas


# ----------------------------------------------------------------------------
# CAMPANHAS INSTITUCIONAIS
# ----------------------------------------------------------------------------

RE_ABA_CPAM = re.compile(r"^CPAM-?\d+$|^CPC$", re.IGNORECASE)


def _extrair_mes_ano_ref(texto: str, ano_pasta: int | None) -> tuple[int | None, int | None]:
    t = _norm(texto)
    ano = ano_pasta
    m = re.search(r"(20\d{2})", t)
    if m:
        ano = int(m.group(1))
    mes = None
    for i, abv in enumerate(MESES_ABV):
        if abv in t:
            mes = i + 1
            break
    return mes, ano


def carregar_campanhas(descartes: list) -> list[dict]:
    bruto = []
    padrao = os.path.join(RAIZ_UNC, "**", "*ampanha*.xlsx")
    arquivos = [
        f for f in glob.glob(padrao, recursive=True)
        if not eh_arquivo_lixo(os.path.basename(f))
    ]
    arquivos_lidos = 0
    for caminho in arquivos:
        base = os.path.basename(caminho)
        m_ano = re.search(r"[\\/](20\d{2})[\\/]", caminho)
        ano_pasta = int(m_ano.group(1)) if m_ano else None
        try:
            wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
        except Exception as e:  # noqa: BLE001 — arquivo corrompido/protegido, pula
            descartes.append(f"Campanhas: falha ao abrir {base}: {e}")
            continue
        usou_arquivo = False
        for nome_aba in wb.sheetnames:
            if not RE_ABA_CPAM.match(nome_aba.strip()):
                continue
            ws = wb[nome_aba]
            try:
                rows = list(ws.iter_rows(max_col=3, values_only=True))
            except Exception:  # noqa: BLE001
                continue
            if not rows:
                continue
            mes_ref_txt = None
            for r in rows[:3]:
                joined = " ".join(_celula(c) for c in r)
                if "REFER" in _norm(joined):
                    mes_ref_txt = joined
                    break
            mes, ano = _extrair_mes_ano_ref(mes_ref_txt or "", ano_pasta)
            i = 0
            while i < len(rows):
                c0 = _norm(rows[i][0] if len(rows[i]) > 0 else None)
                if c0.startswith("16") and "BPM" in c0:
                    j = i
                    while j < len(rows):
                        r = rows[j]
                        cj = _norm(r[0] if len(r) > 0 else None)
                        if j > i and cj:
                            break
                        qtd = r[1] if len(r) > 1 else None
                        nome_camp = _celula(r[2] if len(r) > 2 else None)
                        if nome_camp:
                            try:
                                qtd_num = int(qtd) if qtd not in (None, "") else None
                            except (TypeError, ValueError):
                                qtd_num = None
                            bruto.append({
                                "ano": ano, "mes": mes, "campanha": nome_camp,
                                "quantidade": qtd_num,
                                "origem_arquivo": base, "aba": nome_aba,
                            })
                            usou_arquivo = True
                        j += 1
                    i = j
                    continue
                i += 1
        wb.close()
        if usou_arquivo:
            arquivos_lidos += 1

    # dedupe: dezenas de cópias do mesmo relatório mensal circulam entre
    # pastas de Cia/CPAM — mantém a maior quantidade observada por
    # (ano, mês, nome da campanha normalizado).
    agregados: dict[tuple, dict] = {}
    sem_mes = 0
    for r in bruto:
        if r["ano"] is None or r["mes"] is None:
            sem_mes += 1
            continue
        chave = (r["ano"], r["mes"], _norm(r["campanha"]))
        atual = agregados.get(chave)
        if atual is None or (r["quantidade"] or 0) > (atual["quantidade"] or 0):
            agregados[chave] = r

    if sem_mes:
        descartes.append(f"Campanhas: {sem_mes} linha(s) brutas sem mês/ano identificável na 'MÊS DE REFERÊNCIA' — descartadas.")
    descartes.append(f"Campanhas: {len(arquivos)} arquivo(s) candidato(s), {arquivos_lidos} com bloco 16º BPM/M real, {len(bruto)} linhas brutas -> {len(agregados)} após dedupe por (ano,mês,campanha).")

    linhas = []
    for (ano, mes, _chave_norm), r in agregados.items():
        linhas.append({
            "ano": ano, "mes": mes, "cia": None, "campanha": r["campanha"],
            "quantidade": r["quantidade"], "unidade": None,
            "origem_arquivo": r["origem_arquivo"], "aba": r["aba"],
        })
        registrar_arquivo_p5("campanhas", os.path.join(RAIZ_UNC, r["origem_arquivo"]), r["aba"], 1, f"{ano}-{mes:02d} {r['campanha']}")
    return linhas


# ----------------------------------------------------------------------------
# GOVERNANÇA — Indicadores dos Objetivos Estratégicos
# ----------------------------------------------------------------------------

def carregar_governanca(descartes: list) -> list[dict]:
    """Só grava agregado_dimensional (secao='governanca', sem CHECK de
    schema). NÃO grava fato_secao — ver docstring do módulo."""
    linhas = []
    caminho = os.path.join(
        RAIZ_UNC, "2025", "Documento de Prazo",
        "TRIMESTREAL JUN Cópia de Indicadores dos Objetivos Estratégicos JUN.xlsx",
    )
    if not os.path.isfile(caminho):
        descartes.append(f"Governança objetivos estratégicos: arquivo não encontrado {caminho}")
        return linhas

    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    aba = "C Com Soc"
    if aba not in wb.sheetnames:
        descartes.append(f"Governança objetivos estratégicos: aba '{aba}' ausente em {caminho}")
        wb.close()
        return linhas

    ws = wb[aba]
    rows = list(ws.iter_rows(values_only=True))
    lidas = 0
    for r in rows[2:]:  # linha 0 = ano, linha 1 = cabeçalho
        objetivo = _celula(r[0] if len(r) > 0 else None)
        if not objetivo:
            continue
        meses_valores = r[2:14]  # colunas C..N = JAN..DEZ
        valor_recente = None
        for v in meses_valores:
            if isinstance(v, (int, float)):
                valor_recente = float(v)
        if valor_recente is None:
            continue
        lidas += 1
        linhas.append({
            "secao": "governanca", "fonte": "objetivos_estrategicos",
            "dimensao": "objetivo_estrategico",
            "chave": objetivo[:500],
            "valor": valor_recente,
        })
    wb.close()
    descartes.append(
        f"Governança: {lidas} objetivo(s) estratégico(s) gravados em agregado_dimensional "
        f"(valor = último mês com dado numérico em 2025, JAN..OUT). fato_secao indicador="
        f"'objetivos_no_alvo' NÃO gravado: fato_secao_secao_check não permite secao='governanca' "
        f"(migration 007 restringe a p1/p2/p3/p4/p5/ft/motomec/res_armas/spjmd) — requer nova "
        f"migration para estender o CHECK, e por instrução do projeto nenhuma alteração de schema "
        f"é feita sem autorização prévia do usuário."
    )
    registrar_arquivo_p5("governanca", caminho, aba, lidas, "Indicadores dos Objetivos Estratégicos 2025, valor = mês mais recente por objetivo")
    return linhas


# ----------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    if not caminho_disponivel(RAIZ_UNC):
        print(f"[p5_reconhecimento] Fonte indisponível: {RAIZ_UNC}")
        if not args.dry_run:
            batch_id = abrir_batch(SECAO, FONTE)
            fechar_batch(batch_id, "fonte_indisponivel", erro="UNC inacessível no momento da execução")
        sys.exit(0)

    descartes: list[str] = []
    resultado = {}

    tabelas_alvo = {
        "p5_lmp": ("lmp", carregar_lmp),
        "p5_agraciados": ("agraciados", carregar_agraciados),
        "p5_indicados": ("indicados", carregar_indicados),
        "p5_campanhas": ("campanhas", carregar_campanhas),
    }

    batch_id = None if args.dry_run else abrir_batch(SECAO, FONTE)
    total_validas = 0
    erro_fatal = None

    try:
        for tabela, (nome, funcao) in tabelas_alvo.items():
            if not args.dry_run and not tabela_vazia(tabela):
                descartes.append(f"{tabela}: já contém linhas — pulado (sem unique constraint além de id, não é seguro reingerir).")
                print(f"[p5_reconhecimento] {tabela}: NÃO vazia, pulando extração/carga.")
                continue
            linhas = funcao(descartes)
            print(f"[p5_reconhecimento] {tabela}: {len(linhas)} linha(s) extraída(s).")
            resultado[tabela] = len(linhas)
            total_validas += len(linhas)
            linhas = normalizar_chaves(linhas)
            if not args.dry_run and linhas:
                for i in range(0, len(linhas), 500):
                    resp = requests.post(
                        _url(tabela),
                        headers={**_HEADERS, "Prefer": "return=minimal"},
                        json=linhas[i:i + 500],
                        timeout=60,
                    )
                    resp.raise_for_status()

        # governança — agregado_dimensional (framework genérico, upsert real)
        linhas_gov = carregar_governanca(descartes)
        print(f"[p5_reconhecimento] agregado_dimensional (governanca): {len(linhas_gov)} linha(s).")
        resultado["agregado_dimensional(governanca)"] = len(linhas_gov)
        total_validas += len(linhas_gov)
        if not args.dry_run and linhas_gov:
            upsert_agregado_dimensional(linhas_gov)

    except Exception as e:  # noqa: BLE001
        erro_fatal = str(e)

    if args.dry_run:
        print(f"[p5_reconhecimento] dry-run concluído: {total_validas} linhas válidas no total.")
        for d in descartes:
            print("  -", d)
        return

    if erro_fatal:
        fechar_batch(batch_id, "falha", linhas_validas=total_validas, descartes=descartes, erro=erro_fatal)
        print(f"[p5_reconhecimento] FALHA: {erro_fatal}")
        sys.exit(1)

    fechar_batch(batch_id, "ok" if total_validas else "parcial", linhas_validas=total_validas, descartes=descartes)
    print(f"[p5_reconhecimento] Concluído: {total_validas} linhas válidas no total.")
    for tabela, n in resultado.items():
        print(f"  - {tabela}: {n}")
    print("Observações/descartes:")
    for d in descartes:
        print("  -", d)


if __name__ == "__main__":
    main()
